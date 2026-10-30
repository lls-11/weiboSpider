# -*- coding: UTF-8 -*-
"""把 markdown 报告转成 docx（标题、段落、列表、表格、行内加粗）。

用法：python data/md_to_docx.py docs/research_report_psych.md docs/研究分析报告_心理学框架.docx
只处理本项目报告用到的语法子集，不做通用转换。
"""

import os
import re
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

BOLD = re.compile(r'\*\*(.+?)\*\*')
IMAGE = re.compile(r'^!\[(.*?)\]\((.+?)\)\s*$')


def add_runs(par, text):
    """把 **加粗** 拆成 run。"""
    pos = 0
    for m in BOLD.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        r = par.add_run(m.group(1))
        r.bold = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    for i, c in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        r = p.add_run(c.strip())
        r.bold = True
        r.font.size = Pt(9)
    for row in body:
        cells = t.add_row().cells
        for i, c in enumerate(row[:len(header)]):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            add_runs(p, c.strip())
            for r in p.runs:
                r.font.size = Pt(9)
    doc.add_paragraph()


def convert(md_path, out_path):
    lines = open(md_path, encoding='utf-8').read().split('\n')
    doc = docx.Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(10.5)
    buf = []
    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.rstrip()
        # 图片
        mi = IMAGE.match(s.strip())
        if mi:
            path = mi.group(2)
            if os.path.exists(path):
                doc.add_picture(path, width=docx.shared.Cm(15.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                if mi.group(1):
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = cap.add_run(mi.group(1))
                    r.font.size = Pt(9)
            else:
                doc.add_paragraph('[缺图：%s]' % path)
            i += 1
            continue
        # 表格块
        if s.startswith('|') and i + 1 < len(lines) and set(lines[i + 1].replace('|', '').replace(' ', '')) <= set('-:'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c for c in lines[i].strip().strip('|').split('|')]
                if not set(''.join(cells).replace(' ', '')) <= set('-:'):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        if s.startswith('#### '):
            doc.add_heading(s[5:].strip(), level=4)
        elif s.startswith('### '):
            doc.add_heading(s[4:].strip(), level=3)
        elif s.startswith('## '):
            doc.add_heading(s[3:].strip(), level=2)
        elif s.startswith('# '):
            h = doc.add_heading(s[2:].strip(), level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif s.startswith('---'):
            pass
        elif s.startswith('*') and s.endswith('*') and not s.startswith('**') and len(s) > 2:
            p = doc.add_paragraph()
            add_runs(p, s.strip('*'))
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        elif re.match(r'^\d+\.\s', s):
            p = doc.add_paragraph(style='List Number')
            add_runs(p, re.sub(r'^\d+\.\s', '', s))
        elif s.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, s[2:])
        elif s.strip() == '':
            if buf:
                p = doc.add_paragraph()
                add_runs(p, ' '.join(buf))
                buf = []
        else:
            buf.append(s.strip())
        i += 1
    if buf:
        p = doc.add_paragraph()
        add_runs(p, ' '.join(buf))
    doc.save(out_path)
    return out_path


if __name__ == '__main__':
    src, dst = sys.argv[1], sys.argv[2]
    os.makedirs(os.path.dirname(dst) or '.', exist_ok=True)
    print('已生成', convert(src, dst))
